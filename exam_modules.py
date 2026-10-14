"""试卷答案查找、重点标记和重复考点分析。

对接说明：本模块只通过 ``LibraryDatabase.search()`` 调用现有关键词检索，
不修改资料导入、分块、索引或搜索排序逻辑。较长题干会拆成若干原文短语，
逐一调用原搜索接口，再从返回结果中选择相关度最高的知识库原文。

存储说明：标记、试卷任务、题目、命中关系和检索日志保存在学习库 SQLite
数据库的独立表中。原始 PDF/Word/Markdown 文件以及原有 documents/chunks 表
均不被改写。
"""

from __future__ import annotations

import json
import re
import time
from dataclasses import dataclass, field, replace
from datetime import datetime
from difflib import SequenceMatcher
from pathlib import Path
from typing import Callable, Iterable

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

from library_core import LibraryDatabase, SearchHit, extract_file, normalize_text


QUESTION_START = re.compile(
    r"^\s*(?:#{1,6}\s*)?(?:(\d{1,4})\s*[\.．、)）:]|第\s*(\d{1,4})\s*题\s*[\.．、:]?)\s*(.*)$"
)
OPTION_START = re.compile(r"^\s*[A-HＡ-Ｈ]\s*[\.．、)）:]\s*", re.I)
PUNCTUATION = re.compile(r"[\s，。；：、！？,.!?;:（）()《》【】\[\]“”‘’_—-]+")
SENTENCE_SPLIT = re.compile(r"(?<=[。！？；!?;])|\n+")
ANSWER_LINE = re.compile(r"^\s*(?:参考)?答案\s*[:：]\s*(.*)$", re.I)
ANALYSIS_LINE = re.compile(r"^\s*(?:解析|说明|解答)\s*[:：]", re.I)
STOP_PHRASES = {
    "下列说法正确的是", "下列说法错误的是", "以下说法正确的是", "以下说法错误的是",
    "根据有关规定", "依据有关规定", "应当如何处理", "应该如何处理", "下列哪一项",
    "下列哪些", "请选择", "正确答案", "参考答案", "本题", "关于",
}
XML_INVALID_CHARACTERS = re.compile(
    "[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x84\x86-\x9f\ud800-\udfff\ufffe\uffff]"
)


def clean_export_text(value: object) -> str:
    """Remove invisible characters rejected by Word/XML while preserving text and line breaks."""
    return XML_INVALID_CHARACTERS.sub("", "" if value is None else str(value))


@dataclass
class ExamQuestion:
    number: str
    stem: str
    raw_text: str
    location: str
    order_no: int
    id: int = 0


@dataclass
class ExamMatch:
    question_id: int
    question_number: str
    chunk_id: int
    doc_id: int
    source_name: str
    source_path: str
    source_location: str
    original_text: str
    similarity: float
    query_terms: list[str] = field(default_factory=list)
    marker_types: list[str] = field(default_factory=list)
    id: int = 0
    # Immutable source-chunk snapshot; answer text remains the selected verbatim sentence.
    source_text: str = ""


@dataclass
class ExamTaskResult:
    task_id: int
    task_name: str
    source_path: str
    created_at: str
    questions: list[ExamQuestion]
    matches: dict[int, ExamMatch]
    logs: list[str]


@dataclass
class Marker:
    id: int
    chunk_id: int
    doc_id: int
    document_name: str
    document_path: str
    location: str
    selected_text: str
    start_offset: int
    end_offset: int
    marker_type: str
    note: str
    created_at: str


def _clean_question_stem(raw_text: str) -> str:
    lines = [line.strip() for line in raw_text.splitlines() if line.strip()]
    stem_lines: list[str] = []
    for line in lines:
        if OPTION_START.match(line):
            break
        stem_lines.append(line)
    stem = normalize_text(" ".join(stem_lines))
    stem = re.sub(r"\s*[A-HＡ-Ｈ]\s*[\.．、)）:]\s*.*$", "", stem)
    return stem.strip()


def parse_exam(path: str | Path) -> list[ExamQuestion]:
    """Parse numbered questions from an existing supported document extractor."""
    source = Path(path)
    chunks = extract_file(source)
    questions: list[ExamQuestion] = []
    current_number = ""
    current_location = ""
    buffer: list[str] = []

    def flush() -> None:
        nonlocal buffer, current_number, current_location
        if not buffer:
            return
        raw = normalize_text("\n".join(buffer))
        stem = _clean_question_stem(raw)
        if stem:
            questions.append(
                ExamQuestion(current_number or str(len(questions) + 1), stem, raw, current_location, len(questions) + 1)
            )
        buffer = []

    for chunk in chunks:
        for line in chunk.content.splitlines():
            line = line.strip()
            if not line:
                continue
            match = QUESTION_START.match(line)
            if match:
                flush()
                current_number = match.group(1) or match.group(2) or str(len(questions) + 1)
                current_location = chunk.location
                remainder = match.group(3).strip()
                buffer = [remainder] if remainder else []
            elif current_number:
                buffer.append(line)
        # A Word paragraph or PDF page can itself be one unnumbered question.
        if not current_number and chunk.content.strip().endswith(("？", "?")):
            current_number = str(len(questions) + 1)
            current_location = chunk.location
            buffer = [chunk.content]
            flush()
            current_number = ""
    flush()
    if not questions:
        raise ValueError("未识别到带题号的试题，请确认题目以“1.”、“1、”或“第1题”开头")
    return questions


def _compact(text: str) -> str:
    return PUNCTUATION.sub("", text).casefold()


def _ngrams(text: str, size: int = 2) -> set[str]:
    compact = _compact(text)
    if len(compact) < size:
        return {compact} if compact else set()
    return {compact[index:index + size] for index in range(len(compact) - size + 1)}


def text_similarity(left: str, right: str) -> float:
    """Character-level similarity suited to Chinese source fragments."""
    a, b = _compact(left), _compact(right)
    if not a or not b:
        return 0.0
    grams_a, grams_b = _ngrams(a), _ngrams(b)
    coverage = len(grams_a & grams_b) / max(1, len(grams_a))
    sequence = SequenceMatcher(None, a, b).ratio()
    contains_bonus = 1.0 if a in b or b in a else 0.0
    return min(1.0, coverage * 0.62 + sequence * 0.28 + contains_bonus * 0.10)


def make_query_terms(stem: str, maximum: int = 8) -> list[str]:
    """Build 2-25 character literal phrases for the keyword-search API."""
    cleaned = stem
    for phrase in STOP_PHRASES:
        cleaned = cleaned.replace(phrase, "")
    clauses = [PUNCTUATION.sub("", item) for item in re.split(r"[，。；：！？,.!?;:]", cleaned)]
    candidates: list[str] = []
    for clause in clauses:
        if len(clause) < 2:
            continue
        if len(clause) <= 25:
            candidates.append(clause)
        else:
            window = 12
            positions = (0, max(0, (len(clause) - window) // 2), max(0, len(clause) - window))
            candidates.extend(clause[pos:pos + window] for pos in positions)
    compact = _compact(cleaned)
    if not candidates and compact:
        candidates.append(compact[:25])
    unique = list(dict.fromkeys(item for item in candidates if 2 <= len(item) <= 25))
    unique.sort(key=lambda item: (-len(item), item))
    return unique[:maximum]


def split_exam_content(text: str) -> dict[str, str]:
    """Split source text into visible modules without generating missing content."""
    lines = [line.strip() for line in text.replace("\r", "\n").split("\n") if line.strip()]
    sections: dict[str, list[str]] = {"question": [], "options": [], "answer": [], "analysis": []}
    section = "question"
    option_pattern = re.compile(r"^[A-HＡ-Ｈ][\.．、\)）:]\s*")
    answer_pattern = re.compile(r"^(?:参考)?答案\s*[:：]?", re.I)
    analysis_pattern = re.compile(r"^(?:解析|说明|依据|解答)\s*[:：]?", re.I)
    for line in lines:
        if answer_pattern.match(line):
            section = "answer"
            sections[section].append(answer_pattern.sub("", line).strip())
        elif analysis_pattern.match(line):
            section = "analysis"
            sections[section].append(analysis_pattern.sub("", line).strip())
        elif option_pattern.match(line) and section in {"question", "options"}:
            section = "options"
            sections[section].append(line)
        else:
            sections[section].append(line)
    combined = "\n".join(lines)
    answer_text = "\n".join(item for item in sections["answer"] if item)
    if sections["options"]:
        letters = re.findall(r"\b[A-H]\b", answer_text.upper())
        question_type = "多选题" if len(set(letters)) > 1 else "单选题"
    elif re.search(r"计算|求|公式|=|[＋+－\-×÷/]", combined):
        question_type = "计算题"
    elif re.search(r"判断|正确|错误|对错", combined) and len(combined) < 300:
        question_type = "判断题"
    elif re.search(r"填空|____|（\s*）", combined):
        question_type = "填空题"
    elif sections["question"]:
        question_type = "简答题"
    else:
        question_type = "题型待确认"
    return {
        "type": question_type,
        "question": "\n".join(sections["question"]),
        "options": "\n".join(sections["options"]),
        "answer": answer_text,
        "analysis": "\n".join(item for item in sections["analysis"] if item),
    }


def _best_original_sentence(content: str, stem: str, terms: Iterable[str]) -> str:
    # Prefer an explicitly labelled answer block and stop before analysis or the
    # next numbered question. This keeps only the answer instead of exporting a
    # long surrounding chapter.
    lines = [line.strip() for line in content.replace("\r", "\n").split("\n")]
    for index, line in enumerate(lines):
        marker = ANSWER_LINE.match(line)
        if not marker:
            continue
        answer_lines: list[str] = []
        if marker.group(1).strip():
            answer_lines.append(marker.group(1).strip())
        for following in lines[index + 1:]:
            if not following:
                if answer_lines:
                    break
                continue
            if QUESTION_START.match(following) or ANALYSIS_LINE.match(following):
                break
            answer_lines.append(following)
        if answer_lines:
            return "\n".join(answer_lines)

    sentences = [item.strip() for item in SENTENCE_SPLIT.split(content) if item.strip()]
    if not sentences:
        return content.strip()
    terms = list(terms)

    # A question line repeated in a question bank is not itself an answer.
    declarative = [item for item in sentences if not item.rstrip().endswith(("？", "?"))]
    if declarative:
        sentences = declarative

    def score(sentence: str) -> tuple[float, int]:
        term_score = sum(len(term) for term in terms if term.casefold() in sentence.casefold())
        return text_similarity(stem, sentence), term_score

    best = max(sentences, key=score)
    # Return only verbatim source text. No completion, rewriting, or inferred answer.
    return best


class ExamService:
    def __init__(self, database: LibraryDatabase) -> None:
        self.database = database
        self.initialize()

    def initialize(self) -> None:
        with self.database.connect() as connection:
            connection.executescript(
                """
                CREATE TABLE IF NOT EXISTS knowledge_markers (
                    id INTEGER PRIMARY KEY,
                    chunk_id INTEGER NOT NULL REFERENCES chunks(id) ON DELETE CASCADE,
                    doc_id INTEGER NOT NULL REFERENCES documents(id) ON DELETE CASCADE,
                    selected_text TEXT NOT NULL,
                    start_offset INTEGER NOT NULL DEFAULT 0,
                    end_offset INTEGER NOT NULL DEFAULT 0,
                    marker_type TEXT NOT NULL DEFAULT '考点',
                    note TEXT NOT NULL DEFAULT '',
                    created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
                );
                CREATE INDEX IF NOT EXISTS markers_chunk_id ON knowledge_markers(chunk_id);

                CREATE TABLE IF NOT EXISTS exam_tasks (
                    id INTEGER PRIMARY KEY,
                    name TEXT NOT NULL,
                    source_path TEXT NOT NULL,
                    created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                    question_count INTEGER NOT NULL DEFAULT 0
                );
                CREATE TABLE IF NOT EXISTS exam_questions (
                    id INTEGER PRIMARY KEY,
                    task_id INTEGER NOT NULL REFERENCES exam_tasks(id) ON DELETE CASCADE,
                    question_no TEXT NOT NULL,
                    stem TEXT NOT NULL,
                    raw_text TEXT NOT NULL,
                    source_location TEXT NOT NULL DEFAULT '',
                    order_no INTEGER NOT NULL,
                    status TEXT NOT NULL DEFAULT 'pending',
                    error TEXT NOT NULL DEFAULT ''
                );
                CREATE INDEX IF NOT EXISTS exam_questions_task_id ON exam_questions(task_id);
                CREATE TABLE IF NOT EXISTS exam_matches (
                    id INTEGER PRIMARY KEY,
                    task_id INTEGER NOT NULL REFERENCES exam_tasks(id) ON DELETE CASCADE,
                    question_id INTEGER NOT NULL REFERENCES exam_questions(id) ON DELETE CASCADE,
                    chunk_id INTEGER NOT NULL REFERENCES chunks(id) ON DELETE CASCADE,
                    doc_id INTEGER NOT NULL REFERENCES documents(id) ON DELETE CASCADE,
                    source_name TEXT NOT NULL,
                    source_path TEXT NOT NULL,
                    source_location TEXT NOT NULL,
                    matched_text TEXT NOT NULL,
                    source_text TEXT NOT NULL DEFAULT '',
                    similarity REAL NOT NULL DEFAULT 0,
                    query_terms TEXT NOT NULL DEFAULT '[]',
                    created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
                );
                CREATE INDEX IF NOT EXISTS exam_matches_task_id ON exam_matches(task_id);
                CREATE TABLE IF NOT EXISTS exam_match_markers (
                    match_id INTEGER NOT NULL REFERENCES exam_matches(id) ON DELETE CASCADE,
                    marker_id INTEGER NOT NULL REFERENCES knowledge_markers(id) ON DELETE CASCADE,
                    PRIMARY KEY(match_id, marker_id)
                );
                CREATE TABLE IF NOT EXISTS exam_logs (
                    id INTEGER PRIMARY KEY,
                    task_id INTEGER NOT NULL REFERENCES exam_tasks(id) ON DELETE CASCADE,
                    question_id INTEGER REFERENCES exam_questions(id) ON DELETE CASCADE,
                    level TEXT NOT NULL DEFAULT 'INFO',
                    message TEXT NOT NULL,
                    created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
                );
                """
            )
            # Upgrade only this module's metadata table, never documents/chunks or source files.
            columns = {row["name"] for row in connection.execute("PRAGMA table_info(exam_matches)")}
            if "source_text" not in columns:
                connection.execute("ALTER TABLE exam_matches ADD COLUMN source_text TEXT NOT NULL DEFAULT ''")

    def add_marker(
        self,
        hit: SearchHit,
        selected_text: str,
        marker_type: str = "考点",
        note: str = "",
        start_offset: int | None = None,
    ) -> int:
        selected_text = selected_text.strip()
        if not selected_text:
            raise ValueError("请选择需要标记的原文")
        start = hit.content.find(selected_text) if start_offset is None else max(0, start_offset)
        start = max(0, start)
        end = start + len(selected_text)
        marker_type = marker_type.strip() or "考点"
        with self.database.connect() as connection:
            cursor = connection.execute(
                """INSERT INTO knowledge_markers(
                       chunk_id, doc_id, selected_text, start_offset, end_offset, marker_type, note
                   ) VALUES (?, ?, ?, ?, ?, ?, ?)""",
                (hit.chunk_id, hit.doc_id, selected_text, start, end, marker_type, note.strip()),
            )
            marker_id = int(cursor.lastrowid)
            connection.execute(
                """INSERT OR IGNORE INTO exam_match_markers(match_id, marker_id)
                   SELECT id, ? FROM exam_matches WHERE chunk_id=?""",
                (marker_id, hit.chunk_id),
            )
        return marker_id

    def delete_marker(self, marker_id: int) -> None:
        with self.database.connect() as connection:
            connection.execute("DELETE FROM knowledge_markers WHERE id=?", (marker_id,))

    def list_markers(self) -> list[Marker]:
        with self.database.connect() as connection:
            rows = connection.execute(
                """SELECT m.*, d.name, d.path, c.location
                   FROM knowledge_markers m
                   JOIN documents d ON d.id=m.doc_id
                   JOIN chunks c ON c.id=m.chunk_id
                   ORDER BY m.created_at DESC, m.id DESC"""
            ).fetchall()
        return [
            Marker(
                id=int(row["id"]), chunk_id=int(row["chunk_id"]), doc_id=int(row["doc_id"]),
                document_name=str(row["name"]), document_path=str(row["path"]),
                location=str(row["location"]), selected_text=str(row["selected_text"]),
                start_offset=int(row["start_offset"]), end_offset=int(row["end_offset"]),
                marker_type=str(row["marker_type"]), note=str(row["note"]),
                created_at=str(row["created_at"]),
            )
            for row in rows
        ]

    def _log(self, task_id: int, question_id: int | None, level: str, message: str) -> None:
        with self.database.connect() as connection:
            connection.execute(
                "INSERT INTO exam_logs(task_id, question_id, level, message) VALUES (?, ?, ?, ?)",
                (task_id, question_id, level, message),
            )

    def _context_for_hit(self, hit: SearchHit) -> str:
        rows = self.database.chunk_context(hit.chunk_id, following=3)
        accepted: list[str] = []
        for index, row in enumerate(rows):
            content = str(row["content"]).strip()
            if index and QUESTION_START.match(content):
                break
            accepted.append(content)
            if index and re.search(r"^(?:参考)?答案\s*[:：]|^解析\s*[:：]", content):
                continue
        return "\n\n".join(item for item in accepted if item) or hit.content

    def _create_task(self, source: Path, questions: list[ExamQuestion]) -> tuple[int, list[ExamQuestion]]:
        with self.database.connect() as connection:
            cursor = connection.execute(
                "INSERT INTO exam_tasks(name, source_path, question_count) VALUES (?, ?, ?)",
                (source.stem, str(source.resolve()), len(questions)),
            )
            task_id = int(cursor.lastrowid)
            for question in questions:
                q_cursor = connection.execute(
                    """INSERT INTO exam_questions(
                           task_id, question_no, stem, raw_text, source_location, order_no
                       ) VALUES (?, ?, ?, ?, ?, ?)""",
                    (task_id, question.number, question.stem, question.raw_text, question.location, question.order_no),
                )
                question.id = int(q_cursor.lastrowid)
        return task_id, questions

    def search_question(
        self, question: ExamQuestion, timeout_seconds: float = 8.0,
        activity: Callable[[str], None] | None = None,
        categories: Iterable[str] | None = None,
    ) -> tuple[ExamMatch | None, list[str]]:
        """Call the unchanged keyword interface and return only a verbatim source fragment."""
        started = time.monotonic()
        terms = make_query_terms(question.stem)
        logs: list[str] = [f"检索词：{'；'.join(terms) if terms else '无'}"]
        candidates: dict[int, tuple[SearchHit, set[str]]] = {}
        for term in terms:
            if time.monotonic() - started > timeout_seconds:
                raise TimeoutError(f"检索超过 {timeout_seconds:g} 秒")
            if activity:
                activity(f"第{question.number}题，正在搜索关键词：{term}")
            hits = self.database.search(term, category=list(categories or []), limit=200)
            logs.append(f"“{term}”命中 {len(hits)} 条")
            if activity:
                activity(f"第{question.number}题，“{term}”命中 {len(hits)} 条")
            for hit in hits:
                if hit.chunk_id not in candidates:
                    candidates[hit.chunk_id] = (hit, set())
                candidates[hit.chunk_id][1].add(term)
        if not candidates:
            return None, logs

        ranked: list[tuple[float, int, SearchHit, list[str]]] = []
        question_numbers = set(re.findall(r"\d+(?:\.\d+)?", question.stem))
        for hit, hit_terms in candidates.values():
            relevance = text_similarity(question.stem, hit.content)
            term_weight = sum(len(term) for term in hit_terms) / max(1, sum(len(term) for term in terms))
            exact = 1.0 if _compact(question.stem) in _compact(hit.content) else 0.0
            content_numbers = set(re.findall(r"\d+(?:\.\d+)?", hit.content))
            number_score = 1.0 if not question_numbers else len(question_numbers & content_numbers) / len(question_numbers)
            option_needed = bool(re.search(r"[A-D][\.．、\)]", question.raw_text))
            option_score = 1.0 if not option_needed or re.search(r"[A-D][\.．、\)]", hit.content) else 0.0
            score = min(
                1.0,
                exact * 0.40 + relevance * 0.25 + min(1.0, term_weight) * 0.20
                + number_score * 0.10 + option_score * 0.05,
            )
            if question_numbers and content_numbers and not (question_numbers & content_numbers):
                score *= 0.65
            ranked.append((score, len(hit_terms), hit, sorted(hit_terms)))
        score, _term_count, hit, matched_terms = max(ranked, key=lambda item: (item[0], item[1], item[2].score))
        if score < 0.12:
            logs.append(f"最高相关度 {score:.1%}，低于最低命中条件")
            return None, logs
        context = self._context_for_hit(hit)
        original = _best_original_sentence(context, question.stem, matched_terms)
        explicit_answer = any(ANSWER_LINE.match(line) for line in context.splitlines())
        direct_similarity = text_similarity(question.stem, original)
        compact_original = _compact(original)
        direct_term = any(_compact(term) in compact_original for term in matched_terms if _compact(term))
        if not original.strip() or (
            not explicit_answer and not direct_term and direct_similarity < 0.16
        ):
            logs.append("检索片段与题干没有可直接作答的有效原文，按未命中处理")
            return None, logs
        logs.append(f"采用原文：{hit.name} · {hit.location} · 相关度 {score:.1%}")
        return ExamMatch(
            question_id=question.id, question_number=question.number,
            chunk_id=hit.chunk_id, doc_id=hit.doc_id, source_name=hit.name,
            source_path=hit.path, source_location=hit.location, original_text=original,
            similarity=score, query_terms=matched_terms, source_text=context,
        ), logs

    def search_text_candidates(
        self, query: str, categories: Iterable[str] | None = None, limit: int = 30,
    ) -> list[SearchHit]:
        """Rank multiple candidate questions for interactive trial-paper search."""
        query = query.strip()
        terms = make_query_terms(query)
        if query and len(query) <= 500:
            terms = [query, *terms]
        candidates: dict[int, tuple[SearchHit, set[str]]] = {}
        for term in dict.fromkeys(terms):
            for hit in self.database.search(term, category=list(categories or []), limit=300):
                if hit.chunk_id not in candidates:
                    candidates[hit.chunk_id] = (hit, set())
                candidates[hit.chunk_id][1].add(term)
        query_compact = re.sub(r"\W+", "", query).casefold()
        query_numbers = set(re.findall(r"\d+(?:\.\d+)?", query))
        ranked: list[tuple[float, SearchHit]] = []
        for hit, matched in candidates.values():
            content_compact = re.sub(r"\W+", "", hit.content).casefold()
            exact = 1.0 if query_compact and query_compact in content_compact else 0.0
            similarity = text_similarity(query, hit.content)
            denominator = max(1, sum(len(term) for term in terms if term != query))
            coverage = sum(len(term) for term in matched if term != query) / denominator
            content_numbers = set(re.findall(r"\d+(?:\.\d+)?", hit.content))
            number_score = 1.0 if not query_numbers else len(query_numbers & content_numbers) / len(query_numbers)
            option_needed = bool(re.search(r"[A-D][\.．、\)]", query))
            option_score = 1.0 if not option_needed or re.search(r"[A-D][\.．、\)]", hit.content) else 0.0
            score = exact * 0.40 + similarity * 0.25 + min(1.0, coverage) * 0.20
            score += number_score * 0.10 + option_score * 0.05
            if query_numbers and content_numbers and not (query_numbers & content_numbers):
                score *= 0.65
            hit.score = int(round(min(1.0, score) * 100))
            ranked.append((score, hit))
        ranked.sort(key=lambda item: (-item[0], item[1].name.casefold(), item[1].location))
        results = [hit for score, hit in ranked[:limit] if score >= 0.12]
        for hit in results:
            hit.content = self._context_for_hit(hit)
        return results

    def process_exam(
        self,
        source_path: str | Path,
        progress: Callable[[int, int, ExamQuestion, ExamMatch | None, str], None] | None = None,
        timeout_seconds: float = 8.0,
        activity: Callable[[str], None] | None = None,
        questions: list[ExamQuestion] | None = None,
        categories: Iterable[str] | None = None,
    ) -> ExamTaskResult:
        source = Path(source_path)
        questions = parse_exam(source) if questions is None else questions
        task_id, questions = self._create_task(source, questions)
        matches: dict[int, ExamMatch] = {}
        log_messages: list[str] = []
        self._log(task_id, None, "INFO", f"解析完成，共 {len(questions)} 道题")
        for index, question in enumerate(questions, 1):
            if activity:
                activity(f"开始查找第{question.number}题（{index}/{len(questions)}）")
            status = "未找到知识库原文"
            try:
                match, logs = self.search_question(question, timeout_seconds, activity, categories)
                for item in logs:
                    self._log(task_id, question.id, "INFO", item)
                    log_messages.append(f"第{question.number}题：{item}")
                if match:
                    with self.database.connect() as connection:
                        cursor = connection.execute(
                            """INSERT INTO exam_matches(
                                   task_id, question_id, chunk_id, doc_id, source_name, source_path,
                                   source_location, matched_text, similarity, query_terms, source_text
                               ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                            (
                                task_id, question.id, match.chunk_id, match.doc_id, match.source_name,
                                match.source_path, match.source_location, match.original_text,
                                match.similarity, json.dumps(match.query_terms, ensure_ascii=False), match.source_text,
                            ),
                        )
                        match.id = int(cursor.lastrowid)
                        marker_rows = connection.execute(
                            "SELECT id, marker_type FROM knowledge_markers WHERE chunk_id=?", (match.chunk_id,)
                        ).fetchall()
                        match.marker_types = list(dict.fromkeys(str(row["marker_type"]) for row in marker_rows))
                        connection.executemany(
                            "INSERT OR IGNORE INTO exam_match_markers(match_id, marker_id) VALUES (?, ?)",
                            [(match.id, int(row["id"])) for row in marker_rows],
                        )
                        connection.execute(
                            "UPDATE exam_questions SET status='matched' WHERE id=?", (question.id,)
                        )
                    matches[question.id] = match
                    status = f"命中：{match.source_name}"
                else:
                    with self.database.connect() as connection:
                        connection.execute(
                            "UPDATE exam_questions SET status='unmatched' WHERE id=?", (question.id,)
                        )
            except TimeoutError as exc:
                status = str(exc)
                self._log(task_id, question.id, "TIMEOUT", status)
                log_messages.append(f"第{question.number}题：{status}")
                with self.database.connect() as connection:
                    connection.execute(
                        "UPDATE exam_questions SET status='timeout', error=? WHERE id=?", (status, question.id)
                    )
            except Exception as exc:
                status = f"检索异常：{exc}"
                self._log(task_id, question.id, "ERROR", status)
                log_messages.append(f"第{question.number}题：{status}")
                with self.database.connect() as connection:
                    connection.execute(
                        "UPDATE exam_questions SET status='error', error=? WHERE id=?", (str(exc), question.id)
                    )
            if progress:
                progress(index, len(questions), question, matches.get(question.id), status)
        return ExamTaskResult(
            task_id=task_id, task_name=source.stem, source_path=str(source.resolve()),
            created_at=datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            questions=questions, matches=matches, logs=log_messages,
        )

    def task_markers_for_match(self, match: ExamMatch) -> list[Marker]:
        return [marker for marker in self.list_markers() if marker.chunk_id == match.chunk_id]


def rank_answer_sources(result: ExamTaskResult) -> list[dict]:
    """Count actual adopted answers per source location, NOT keyword hits or fuzzy clusters.

    Multiple keywords for one question count once. Different sentences selected from
    the same source chunk count toward that chunk. Equal text at different locations
    remains separate. The source snapshot is returned verbatim (PDF chunks may be pages).
    """
    groups: dict[tuple, dict] = {}
    for question in result.questions:
        match = result.matches.get(question.id)
        if not match:
            continue
        snapshot = match.source_text or match.original_text
        key = (match.doc_id, match.chunk_id, match.source_location, snapshot)
        group = groups.setdefault(key, {
            "text": snapshot, "source": match.source_name, "location": match.source_location,
            "doc_id": match.doc_id, "chunk_id": match.chunk_id,
            "question_numbers": [], "question_ids": set(), "marker_types": [], "citations": [],
        })
        if question.id in group["question_ids"]:
            continue
        group["question_ids"].add(question.id)
        group["question_numbers"].append(question.number)
        group["citations"].append((question.number, match.original_text))
        group["marker_types"] = list(dict.fromkeys(group["marker_types"] + match.marker_types))
    ranked = list(groups.values())
    for group in ranked:
        group["count"] = len(group["question_ids"])
    return sorted(ranked, key=lambda group: (-group["count"], group["source"], group["location"], group["chunk_id"]))


def analyze_repeated_points(result: ExamTaskResult, threshold: float = 0.78) -> tuple[list[dict], list[dict]]:
    """Group matched verbatim fragments using an adjustable text-similarity threshold."""
    threshold = min(1.0, max(0.0, threshold))
    groups: list[dict] = []
    for question in result.questions:
        match = result.matches.get(question.id)
        if not match:
            continue
        target = None
        for group in groups:
            same_source = group["doc_id"] == match.doc_id
            if same_source and text_similarity(group["text"], match.original_text) >= threshold:
                target = group
                break
        if target is None:
            target = {
                "text": match.original_text, "source": match.source_name,
                "location": match.source_location, "doc_id": match.doc_id,
                "question_numbers": [], "question_ids": set(), "marker_types": list(match.marker_types),
            }
            groups.append(target)
        if question.id not in target["question_ids"]:
            target["question_ids"].add(question.id)
            target["question_numbers"].append(question.number)
        target["marker_types"] = list(dict.fromkeys(target["marker_types"] + match.marker_types))
    repeated = [group for group in groups if len(group["question_numbers"]) > 1]
    single = [group for group in groups if len(group["question_numbers"]) == 1]
    repeated.sort(key=lambda item: (-len(item["question_numbers"]), item["source"]))
    return repeated, single


def combine_exam_results(results: Iterable[ExamTaskResult]) -> ExamTaskResult:
    """Create a read-only combined view for cross-paper repeated-point analysis."""
    items = list(results)
    if not items:
        raise ValueError("没有可合并的试卷结果")
    questions: list[ExamQuestion] = []
    matches: dict[int, ExamMatch] = {}
    logs: list[str] = []
    for paper_index, result in enumerate(items, 1):
        for question in result.questions:
            # Task names and question numbers need not be unique across imported papers.
            name = result.task_name
            if sum(item.task_name == name for item in items) > 1:
                name = f"{name}（卷{paper_index}）"
            label = f"{name}·第{question.number}题"
            combined_id = len(questions) + 1
            questions.append(replace(question, id=combined_id, number=label, order_no=combined_id))
            match = result.matches.get(question.id)
            if match:
                matches[combined_id] = replace(match, question_id=combined_id, question_number=label)
        logs.extend(f"{result.task_name}：{line}" for line in result.logs)
    return ExamTaskResult(
        task_id=0,
        task_name=f"{len(items)}套试卷综合",
        source_path=f"{len(items)}套试卷",
        created_at=datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        questions=questions,
        matches=matches,
        logs=logs,
    )


def _question_label(number: str) -> str:
    return number if "·第" in number or number.endswith("题") else f"第{number}题"


def answer_source_text(match: ExamMatch) -> str:
    """Return only the selected relevant original sentence, never the full context block."""
    return clean_export_text(match.original_text)


def _answer_question_text(question: ExamQuestion) -> str:
    """Use the original question wording, without XML-incompatible hidden characters."""
    original = question.raw_text.strip() or question.stem.strip()
    return clean_export_text(normalize_text(original))


def answer_source_label(match: ExamMatch | None) -> str:
    if match is None:
        return "【来源：未检索到匹配文档】"
    source_name = clean_export_text(match.source_name).strip() or "文件名未标注"
    source_location = clean_export_text(match.source_location).strip() or "位置未标注"
    return f"【来源：{source_name} · {source_location}】"


def answer_markdown(result: ExamTaskResult) -> str:
    lines: list[str] = []
    for question in result.questions:
        match = result.matches.get(question.id)
        lines.extend(["---", f"### 题目：{_answer_question_text(question)}", "", "**答案：**"])
        if match:
            lines.extend([answer_source_text(match), ""])
        else:
            lines.extend(["**【无】**", ""])
        lines.extend(["**出处：**", answer_source_label(match), "---", ""])
    return "\n".join(lines).rstrip() + "\n"


def export_answer_markdown(result: ExamTaskResult, target: str | Path) -> Path:
    path = Path(target)
    path.write_text(answer_markdown(result), encoding="utf-8")
    return path


def apply_exam_word_format(document) -> None:
    """Apply the requested Chinese/Latin typefaces and five-point single spacing."""
    for style in document.styles:
        if getattr(style, "font", None) is not None:
            style.font.name = "Times New Roman"
            style.font.size = Pt(10.5)
            try:
                style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
                style.element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
                style.element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
            except AttributeError:
                pass
    paragraphs = list(document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    for paragraph in paragraphs:
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(10.5)
            run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
            run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
            run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")


def _add_answer_word_group(document, question: ExamQuestion, match: ExamMatch | None) -> None:
    """Write one question using the same fixed question-answer-source structure as Markdown."""
    question_paragraph = document.add_paragraph()
    question_run = question_paragraph.add_run(f"题目：{_answer_question_text(question)}")
    question_run.bold = True

    answer_label = document.add_paragraph()
    answer_label.add_run("答案：").bold = True
    if match:
        document.add_paragraph(answer_source_text(match))
    else:
        empty_paragraph = document.add_paragraph()
        empty_paragraph.add_run("【无】").bold = True

    source_label = document.add_paragraph()
    source_label.add_run("出处：").bold = True
    document.add_paragraph(answer_source_label(match))
    document.add_paragraph("")


def export_answer_word(result: ExamTaskResult, target: str | Path) -> Path:
    document = Document()
    for question in result.questions:
        match = result.matches.get(question.id)
        _add_answer_word_group(document, question, match)
    apply_exam_word_format(document)
    path = Path(target)
    document.save(path)
    return path


def analysis_markdown(result: ExamTaskResult, threshold: float = 0.78) -> str:
    ranked = rank_answer_sources(result)
    repeated = [item for item in ranked if item["count"] > 1]
    single = [item for item in ranked if item["count"] == 1]
    similar, _ = analyze_repeated_points(result, threshold)
    all_matches = [result.matches[q.id] for q in result.questions if q.id in result.matches]
    lines = [
        f"# {result.task_name}——答案出处引用频次与重复考点分析", "",
        "> 按本次试卷答案实际采用的知识库出处统计。每道题对同一出处计 1 次；不是关键词搜索次数，也不是文档中关键词出现次数。", "",
        "> 同一出处内的不同答案句子归入同一原文片段；不同文档或不同位置不合并。来源粒度沿用现有知识库：Word 段落/表格、Markdown 片段、PDF 页/片段，不重新分段，不生成考点。", "",
        "## 一、本次试卷基础信息", "",
        f"- 试卷文件：{Path(result.source_path).name}",
        f"- 建立时间：{result.created_at}",
        f"- 题目数量：{len(result.questions)}",
        f"- 命中知识库原文：{len(all_matches)} 题",
        f"- 命中出处数量：{len(ranked)} 处",
        f"- 最高引用次数：{ranked[0]['count'] if ranked else 0} 次",
        f"- 相似原句辅助分组阈值：{threshold:.0%}（不影响出处排名）", "",
        "## 二、答案出处排行榜（按引用题数从高到低）", "",
    ]
    if not all_matches:
        lines.extend(["暂无考点数据。", ""])
    else:
        for index, item in enumerate(ranked, 1):
            marker = "、".join(item["marker_types"]) or "未标记"
            numbers = "、".join(_question_label(number) for number in item["question_numbers"])
            lines.extend([
                f"### 排名 {index}：被 {item['count']} 道题引用", "",
                f"- 来源：{item['source']} · {item['location']}",
                f"- 对应试卷及题号：{numbers}", f"- 标记：{marker}", "",
                "#### 出处原文片段", "", item["text"], "", "#### 实际采用的答案原句", "",
            ])
            for number, original in item["citations"]:
                lines.extend([f"**{_question_label(number)}**", "", original, ""])
    lines.extend(["## 三、多道题重复引用的出处", ""])
    if repeated:
        for index, item in enumerate(repeated, 1):
            marker = "、".join(item["marker_types"]) if item["marker_types"] else "未标记"
            numbers = "、".join(_question_label(number) for number in item["question_numbers"])
            lines.extend([
                f"### {index}. 重复 {item['count']} 次（{item['count']} 道题引用）", "", item["text"], "",
                f"- 来源：{item['source']} · {item['location']}",
                f"- 对应题号：{numbers}", f"- 标记：{marker}", "",
            ])
    else:
        lines.extend(["暂无被多道题重复引用的出处。", ""])
    lines.extend(["## 四、仅被一道题引用的出处", ""])
    if single:
        for item in single:
            number = item["question_numbers"][0]
            marker = "、".join(item["marker_types"]) if item["marker_types"] else "未标记"
            lines.extend([
                f"### {_question_label(number)}", "", item["text"], "",
                f"- 来源：{item['source']} · {item['location']}", f"- 标记：{marker}", "",
            ])
    else:
        lines.extend(["暂无仅被一道题引用的出处。", ""])
    lines.extend(["## 五、相似答案原句辅助分组", "",
                  "以下仅按相似度辅助对照同一文档内的答案原句，不代表同一段落，不参与上方出处频次排名。", ""])
    if similar:
        for item in similar:
            numbers = "、".join(_question_label(number) for number in item["question_numbers"])
            lines.extend([f"### 相似原句关联 {len(item['question_numbers'])} 道题", "", item["text"], "",
                          f"- 代表原句来源：{item['source']} · {item['location']}", f"- 对应题号：{numbers}", ""])
    else:
        lines.extend(["暂无相似答案原句分组。", ""])
    return "\n".join(lines).rstrip() + "\n"


def export_analysis_markdown(result: ExamTaskResult, target: str | Path, threshold: float = 0.78) -> Path:
    path = Path(target)
    path.write_text(analysis_markdown(result, threshold), encoding="utf-8")
    return path


def batch_answer_markdown(results: Iterable[ExamTaskResult]) -> str:
    lines: list[str] = []
    for result in results:
        for question in result.questions:
            match = result.matches.get(question.id)
            lines.extend(["---", f"### 题目：{_answer_question_text(question)}", "", "**答案：**"])
            if match:
                lines.extend([answer_source_text(match), ""])
            else:
                lines.extend(["**【无】**", ""])
            lines.extend(["**出处：**", answer_source_label(match), "---", ""])
    return "\n".join(lines).rstrip() + "\n"


def export_batch_answer_markdown(results: Iterable[ExamTaskResult], target: str | Path) -> Path:
    path = Path(target)
    path.write_text(batch_answer_markdown(results), encoding="utf-8")
    return path


def export_batch_answer_word(results: Iterable[ExamTaskResult], target: str | Path) -> Path:
    document = Document()
    for result in results:
        for question in result.questions:
            match = result.matches.get(question.id)
            _add_answer_word_group(document, question, match)
    apply_exam_word_format(document)
    path = Path(target)
    document.save(path)
    return path


def export_batch_analysis_markdown(
    results: Iterable[ExamTaskResult], target: str | Path, threshold: float = 0.78
) -> Path:
    items = list(results)
    combined = combine_exam_results(items)
    content = analysis_markdown(combined, threshold)
    source_lines = "\n".join(f"- {item.task_name}：{Path(item.source_path).name}" for item in items)
    content = content.replace(
        f"- 试卷文件：{Path(combined.source_path).name}",
        f"- 试卷数量：{len(items)}\n- 试卷清单：\n{source_lines}",
        1,
    )
    path = Path(target)
    path.write_text(content, encoding="utf-8")
    return path


def _new_export_folder(output_dir: str | Path, timestamp: str | None = None) -> Path:
    """Use an isolated run folder; never overwrite a previous set of answers."""
    root = Path(output_dir)
    root.mkdir(parents=True, exist_ok=True)
    stamp = re.sub(r'[^\w-]', '_', timestamp or datetime.now().strftime('%Y%m%d_%H%M%S_%f'))[:60]
    for index in range(10000):
        folder = root / f"试卷答案_{stamp}{'_' + str(index) if index else ''}"
        try:
            folder.mkdir()
            return folder
        except FileExistsError:
            continue
    raise FileExistsError("无法创建新的答案输出目录，请更换保存位置")


def _paper_file_title(result: ExamTaskResult, index: int) -> str:
    title = re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', result.task_name).strip('. ')[:80] or '试卷'
    # Index prevents collisions for identical names or names normalized to the same string.
    return f"{index:02d}_{title}"


def export_individual_answers(
    results: Iterable[ExamTaskResult], output_dir: str | Path, file_format: str,
) -> tuple[Path, ...]:
    """Manual multi-paper export also writes one answer document per paper."""
    items = list(results)
    if not items:
        raise ValueError("没有可导出的试卷结果")
    if file_format not in ('md', 'docx'):
        raise ValueError("答案格式仅支持 md 或 docx")
    folder = _new_export_folder(output_dir)
    exporter = export_answer_markdown if file_format == 'md' else export_answer_word
    return tuple(exporter(item, folder / f'{_paper_file_title(item, index)}_答案.{file_format}')
                 for index, item in enumerate(items, 1))


def export_result_bundle(
    results: Iterable[ExamTaskResult],
    output_dir: str | Path,
    threshold: float = 0.78,
    timestamp: str | None = None,
) -> tuple[Path, ...]:
    """Generate separate MD/Word answers and analysis for EACH paper, plus cross-paper analysis.

    Ordering is (paper1 MD, Word, analysis, paper2 MD, Word, analysis, ..., batch analysis).
    The batch report never replaces the individual answer documents.
    """
    items = list(results)
    if not items:
        raise ValueError("没有可导出的试卷结果")
    folder = _new_export_folder(output_dir, timestamp)
    outputs: list[Path] = []
    for index, item in enumerate(items, 1):
        title = _paper_file_title(item, index)
        outputs.extend((
            export_answer_markdown(item, folder / f'{title}_答案.md'),
            export_answer_word(item, folder / f'{title}_答案.docx'),
            export_analysis_markdown(item, folder / f'{title}_出处引用分析.md', threshold),
        ))
    if len(items) > 1:
        outputs.append(export_batch_analysis_markdown(items, folder / '全部试卷_跨卷出处引用排名.md', threshold))
    return tuple(outputs)
