import hashlib
import json
from pathlib import Path
from tempfile import TemporaryDirectory

import pymupdf
from docx import Document
from docx.oxml.ns import qn

from exam_modules import (
    ExamMatch,
    ExamQuestion,
    ExamService,
    ExamTaskResult,
    answer_markdown,
    apply_exam_word_format,
    export_answer_word,
    split_exam_content,
)
from library_core import LibraryDatabase
from pdf_viewer import export_annotated_pdf
from document_converter import convert_document_to_searchable_pdf
from question_research import (
    export_question_research_markdown,
    export_question_research_word,
    organize_question_from_library,
)


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def make_pdf(path: Path, text: str) -> None:
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 90), text, fontsize=12)
    document.save(path)
    document.close()


def main() -> None:
    with TemporaryDirectory() as temp:
        root = Path(temp)
        one = root / "制动题库.md"
        two = root / "信号题库.md"
        one.write_text("1. 制动距离由哪些因素决定？\nA. 速度\nB. 制动力\n答案：A", encoding="utf-8")
        two.write_text("1. 信号显示的基本要求是什么？\n答案：保持清晰", encoding="utf-8")
        pdf_source = root / "批注测试.pdf"
        make_pdf(pdf_source, "PDF annotation test source")
        database = LibraryDatabase(root / "library.db", root / "资料文件")
        database.import_file(one, category="制动专业")
        database.import_file(two, category="信号专业")
        database.import_file(pdf_source, category="批注资料")
        rows = database.list_documents()
        first = next(row for row in rows if row["name"] == pdf_source.name)

        assert database.search("制动距离", category=["制动专业"])
        assert not database.search("制动距离", category=["信号专业"])
        service = ExamService(database)
        candidates = service.search_text_candidates("制动距离由哪些因素决定", ["制动专业"])
        assert candidates and candidates[0].name == one.name
        parts = split_exam_content("题目\nA.甲\nB.乙\n答案：A\n解析：依据原文")
        assert parts["options"] and parts["answer"] == "A" and parts["analysis"] == "依据原文"

        original_hash = digest(Path(first["stored_path"]))
        database.add_pdf_annotation(
            int(first["id"]), 0, "highlight", json.dumps({"rect": [60, 60, 260, 105]}), "#ffd54f"
        )
        database.add_pdf_annotation(
            int(first["id"]), 0, "note", json.dumps({"rect": [80, 120, 240, 190]}), "#2196f3", "重点复习"
        )
        annotated = root / "批注版.pdf"
        export_annotated_pdf(database, first, annotated)
        assert annotated.exists() and digest(Path(first["stored_path"])) == original_hash
        check = pymupdf.open(annotated)
        assert sum(1 for _ in (check[0].annots() or [])) == 2
        check.close()

        backup = root / "backup.zip"
        database.export_backup(backup)
        restored = LibraryDatabase(root / "restored.db", root / "恢复资料")
        restored.restore_backup(backup)
        restored_row = restored.list_documents(category="批注资料")[0]
        assert len(restored.list_pdf_annotations(int(restored_row["id"]))) == 2
        renamed = restored.rename_document(int(restored_row["id"]), "批注重点资料.pdf")
        assert renamed["name"] == "批注重点资料.pdf" and Path(renamed["stored_path"]).exists()

        word = Document()
        paragraph = word.add_paragraph("中文 English")
        apply_exam_word_format(word)
        run = paragraph.runs[0]
        assert run.font.size.pt == 10.5
        assert run.font.name == "Times New Roman"
        assert run._element.rPr.rFonts.get(qn("w:eastAsia")) == "宋体"
        assert paragraph.paragraph_format.line_spacing == 1.0

        unsafe_question = ExamQuestion("1", "安全要求？\x00", "安全要求？\x00", "第1段\x0b", 1, id=9001)
        unmatched_question = ExamQuestion("2", "未命中题目？", "未命中题目？\x1f", "第2段", 2, id=9002)
        unsafe_match = ExamMatch(
            9001, "1", 1, 1, "规章\x0c.docx", "x", "第3段\x0b", "原文", 0.9,
            source_text="有效原文\x00保留内容\n第二行\x0b",
        )
        unsafe_result = ExamTaskResult(
            1, "控制字符测试", "test.docx", "2026-09-28",
            [unsafe_question, unmatched_question], {9001: unsafe_match}, [],
        )
        safe_markdown = answer_markdown(unsafe_result)
        assert "原文" in safe_markdown and "有效原文保留内容" not in safe_markdown
        assert "**【无】**" in safe_markdown
        assert "【来源：未检索到匹配文档】" in safe_markdown
        assert not any(character in safe_markdown for character in ("\x00", "\x0b", "\x0c", "\x1f"))
        safe_word = export_answer_word(unsafe_result, root / "控制字符答案.docx")
        safe_word_text = "\n".join(item.text for item in Document(safe_word).paragraphs)
        assert "原文" in safe_word_text and "【无】" in safe_word_text

        source_md = root / "可转换.md"
        source_md.write_text("Markdown中文可检索内容", encoding="utf-8")
        converted_pdf = convert_document_to_searchable_pdf(source_md, root / "可转换.pdf")
        converted = pymupdf.open(converted_pdf)
        assert "Markdown中文可检索内容" in "".join(page.get_text() for page in converted)
        converted.close()

        rule_one = root / "铁运〔2020〕153号规定.md"
        rule_two = root / "接发列车补充规定.md"
        rule_three = root / "铁运〔2022〕88号新规.md"
        original_rule = "第十条 列车接发作业必须核对信号。"
        rule_one.write_text(original_rule, encoding="utf-8")
        rule_two.write_text(original_rule, encoding="utf-8")
        rule_three.write_text("第十条 列车接发作业无需核对信号。", encoding="utf-8")
        database.import_file(rule_one, category="规章")
        database.import_file(rule_two, category="规章")
        database.import_file(rule_three, category="规章")
        research = organize_question_from_library(database, "列车接发作业有哪些规定？", categories="规章")
        assert len(research.statements) == 2, "完全相同的条款应合并为唯一表述"
        same = next(item for item in research.statements if item.text == original_rule)
        assert len(same.sources) == 2, "相同原文的全部来源应保留"
        assert same.conflict_with, "相反条款应同时保留并标注可能冲突"
        rendered = research.render_text()
        assert original_rule in rendered and "可能存在表述冲突" in rendered
        assert "铁运〔2020〕153号" in rendered and "使用的文档名称 / 文号" in rendered
        research_md = export_question_research_markdown(research, root / "问题整理.md")
        research_word = export_question_research_word(research, root / "问题整理.docx")
        assert research_md.exists() and original_rule in research_md.read_text(encoding="utf-8-sig")
        exported = Document(research_word)
        assert any(original_rule in paragraph.text for paragraph in exported.paragraphs)

        temporary_category = database.create_category("待删除分类")
        temporary_doc = root / "待移动资料.md"
        temporary_doc.write_text("分类删除时资料不得丢失。", encoding="utf-8")
        database.import_file(temporary_doc, category=temporary_category)
        assert database.delete_category(temporary_category) == 1
        moved_row = next(row for row in database.list_documents() if row["name"] == temporary_doc.name)
        assert moved_row["category"] == "未分类" and Path(moved_row["stored_path"]).exists()
        assert not (database.storage_dir / temporary_category).exists()

    print("PASS: 分类安全删除、问题整理与双格式导出、试卷候选、PDF批注、备份及PDF转换")


if __name__ == "__main__":
    main()
