from __future__ import annotations

import os
import re
import sqlite3
import subprocess
import hashlib
import json
import shutil
import tempfile
import webbrowser
import zipfile
from contextlib import contextmanager
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Callable, Iterable, Iterator

from docx import Document
from docx.document import Document as DocumentObject
from docx.table import Table
from docx.text.paragraph import Paragraph
from pypdf import PdfReader


SUPPORTED_SUFFIXES = {".pdf", ".docx", ".doc", ".md", ".markdown"}


@dataclass
class TextChunk:
    location: str
    content: str
    order_no: int


@dataclass
class SearchHit:
    chunk_id: int
    doc_id: int
    name: str
    path: str
    location: str
    content: str
    snippet: str
    modified_at: float
    score: int = 1


def normalize_text(text: str) -> str:
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[\t\u00a0]+", " ", text)
    text = re.sub(r"[ ]{2,}", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def normalize_ocr_text(text: str) -> str:
    text = normalize_text(text)
    # Tesseract often inserts a space between every Chinese character. Keeping
    # those spaces would make normal Chinese keyword searches miss the OCR text.
    return re.sub(r"(?<=[\u3400-\u9fff]) +(?=[\u3400-\u9fff])", "", text)


def make_snippet(text: str, keyword: str, radius: int = 180) -> str:
    compact = re.sub(r"\s+", " ", text).strip()
    position = compact.casefold().find(keyword.casefold())
    if position < 0:
        return compact[: radius * 2] + ("…" if len(compact) > radius * 2 else "")
    start = max(0, position - radius)
    end = min(len(compact), position + len(keyword) + radius)
    prefix = "…" if start else ""
    suffix = "…" if end < len(compact) else ""
    return prefix + compact[start:end] + suffix


def _iter_docx_blocks(parent: DocumentObject) -> Iterator[Paragraph | Table]:
    body = parent.element.body
    for child in body.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, parent)
        elif child.tag.endswith("}tbl"):
            yield Table(child, parent)


def extract_docx(path: Path) -> list[TextChunk]:
    document = Document(str(path))
    chunks: list[TextChunk] = []
    paragraph_no = 0
    table_no = 0
    order_no = 0
    for block in _iter_docx_blocks(document):
        order_no += 1
        if isinstance(block, Paragraph):
            paragraph_no += 1
            text = normalize_text(block.text)
            if text:
                chunks.append(TextChunk(f"第 {paragraph_no} 段", text, order_no))
        else:
            table_no += 1
            lines: list[str] = []
            for row in block.rows:
                cells = [normalize_text(cell.text) for cell in row.cells]
                if any(cells):
                    lines.append("\t".join(cells))
            text = "\n".join(lines).strip()
            if text:
                chunks.append(TextChunk(f"表格 {table_no}", text, order_no))
    return chunks


def extract_pdf(path: Path) -> list[TextChunk]:
    reader = PdfReader(str(path))
    chunks: list[TextChunk] = []
    blank_pages: list[int] = []
    ocr_error = ""
    for index, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text(extraction_mode="layout") or ""
        except TypeError:
            text = page.extract_text() or ""
        text = normalize_text(text)
        if text:
            chunks.append(TextChunk(f"第 {index} 页", text, index))
        else:
            blank_pages.append(index)

    # Image-only pages have no PDF text layer. OCR them locally so scanned
    # regulations and examination papers can still be indexed and searched.
    if blank_pages:
        ocr_text, ocr_error = _ocr_pdf_pages(path, blank_pages)
        for page_no in blank_pages:
            text = normalize_ocr_text(ocr_text.get(page_no, ""))
            if text:
                chunks.append(TextChunk(f"第 {page_no} 页（OCR）", text, page_no))
    if not chunks:
        if ocr_error:
            raise ValueError(ocr_error)
        raise ValueError("扫描型 PDF 已尝试自动 OCR，但没有识别到可搜索文字；请检查扫描清晰度")
    chunks.sort(key=lambda chunk: chunk.order_no)
    return chunks


def _find_tesseract() -> tuple[Path | None, Path | None, str]:
    candidates: list[Path] = []
    configured = os.getenv("TESSERACT_CMD", "").strip()
    if configured:
        candidates.append(Path(configured))
    located = shutil.which("tesseract")
    if located:
        candidates.append(Path(located))
    candidates.extend(
        [
            Path(os.environ.get("ProgramFiles", r"C:\Program Files")) / "Tesseract-OCR" / "tesseract.exe",
            Path(os.environ.get("ProgramFiles(x86)", r"C:\Program Files (x86)"))
            / "Tesseract-OCR"
            / "tesseract.exe",
        ]
    )
    executable = next((item for item in candidates if item.is_file()), None)
    if executable is None:
        return None, None, ""
    tessdata = executable.parent / "tessdata"
    available = {item.stem for item in tessdata.glob("*.traineddata")} if tessdata.is_dir() else set()
    if {"chi_sim", "eng"}.issubset(available):
        language = "chi_sim+eng"
    elif "chi_sim" in available:
        language = "chi_sim"
    elif "eng" in available:
        language = "eng"
    else:
        language = ""
    return executable, tessdata if tessdata.is_dir() else None, language


def _ocr_pdf_pages(path: Path, page_numbers: list[int]) -> tuple[dict[int, str], str]:
    executable, tessdata, language = _find_tesseract()
    if executable is None:
        return {}, "该 PDF 为图片扫描件，但未找到本机 OCR 引擎（Tesseract）"
    try:
        import pymupdf
    except ImportError:
        return {}, "该 PDF 为图片扫描件，但软件缺少 PDF OCR 渲染组件"

    results: dict[int, str] = {}
    environment = os.environ.copy()
    if tessdata:
        environment["TESSDATA_PREFIX"] = str(tessdata)
    creation_flags = getattr(subprocess, "CREATE_NO_WINDOW", 0)
    errors: list[str] = []
    try:
        with pymupdf.open(path) as document, tempfile.TemporaryDirectory() as temporary:
            temporary_dir = Path(temporary)
            for page_no in page_numbers:
                image = temporary_dir / f"page-{page_no}.png"
                page = document.load_page(page_no - 1)
                pixmap = page.get_pixmap(dpi=260, colorspace=pymupdf.csGRAY, alpha=False)
                pixmap.save(image)
                command = [str(executable), str(image), "stdout"]
                if language:
                    command.extend(["-l", language])
                command.extend(["--psm", "6"])
                try:
                    completed = subprocess.run(
                        command,
                        capture_output=True,
                        text=True,
                        encoding="utf-8",
                        errors="replace",
                        timeout=180,
                        check=False,
                        env=environment,
                        creationflags=creation_flags,
                    )
                    if completed.returncode == 0:
                        results[page_no] = completed.stdout
                    else:
                        errors.append(completed.stderr.strip() or f"第 {page_no} 页 OCR 失败")
                except subprocess.TimeoutExpired:
                    errors.append(f"第 {page_no} 页 OCR 超时")
    except Exception as exc:
        errors.append(str(exc))
    if results:
        return results, ""
    detail = errors[0] if errors else "没有识别到文字"
    return {}, f"扫描型 PDF 自动 OCR 失败：{detail}"


def _read_text_file(path: Path) -> str:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "utf-16"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def extract_markdown(path: Path) -> list[TextChunk]:
    source = _read_text_file(path).replace("\r\n", "\n").replace("\r", "\n")
    chunks: list[TextChunk] = []
    buffer: list[str] = []
    start_line = 1

    def flush(end_line: int) -> None:
        nonlocal buffer, start_line
        text = normalize_text("\n".join(buffer))
        if text:
            location = f"第 {start_line} 行" if start_line == end_line else f"第 {start_line}-{end_line} 行"
            chunks.append(TextChunk(location, text, len(chunks) + 1))
        buffer = []

    for line_no, line in enumerate(source.split("\n"), start=1):
        if not line.strip():
            if buffer:
                flush(line_no - 1)
            start_line = line_no + 1
        else:
            if not buffer:
                start_line = line_no
            buffer.append(line)
    if buffer:
        flush(len(source.split("\n")))
    return chunks


def extract_legacy_doc(path: Path) -> list[TextChunk]:
    try:
        import pythoncom
        import win32com.client
    except ImportError as exc:
        raise ValueError("读取旧版 .doc 需要安装 Microsoft Word；建议另存为 .docx 后再导入") from exc

    pythoncom.CoInitialize()
    word = None
    document = None
    try:
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        document = word.Documents.Open(str(path.resolve()), ReadOnly=True, AddToRecentFiles=False)
        text = document.Content.Text.replace("\r\x07", "\n").replace("\r", "\n")
        parts = [normalize_text(item) for item in re.split(r"\n{2,}", text)]
        return [TextChunk(f"第 {i} 段", item, i) for i, item in enumerate(parts, 1) if item]
    except Exception as exc:
        raise ValueError("无法读取旧版 .doc；请确认 Microsoft Word 已安装，或先另存为 .docx") from exc
    finally:
        if document is not None:
            document.Close(False)
        if word is not None:
            word.Quit()
        pythoncom.CoUninitialize()


def extract_file(path: Path) -> list[TextChunk]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return extract_pdf(path)
    if suffix == ".docx":
        return extract_docx(path)
    if suffix == ".doc":
        return extract_legacy_doc(path)
    if suffix in {".md", ".markdown"}:
        return extract_markdown(path)
    raise ValueError(f"不支持的文件类型：{suffix}")


def collect_supported_files(paths: Iterable[str | Path]) -> list[Path]:
    found: dict[str, Path] = {}
    for item in paths:
        path = Path(item).expanduser()
        if path.is_file() and path.suffix.lower() in SUPPORTED_SUFFIXES:
            found[str(path.resolve()).casefold()] = path.resolve()
        elif path.is_dir():
            for candidate in path.rglob("*"):
                if (
                    candidate.is_file()
                    and candidate.suffix.lower() in SUPPORTED_SUFFIXES
                    and not any(part.startswith("~$") for part in candidate.parts)
                ):
                    found[str(candidate.resolve()).casefold()] = candidate.resolve()
    return sorted(found.values(), key=lambda value: str(value).casefold())


class LibraryDatabase:
    def __init__(self, db_path: Path, storage_dir: Path | None = None) -> None:
        self.db_path = db_path
        self.storage_dir = storage_dir or db_path.parent / "资料文件"
        self.db_path.parent.mkdir(parents=True, exist_ok=True)
        self.storage_dir.mkdir(parents=True, exist_ok=True)
        self.initialize()

    @contextmanager
    def connect(self) -> Iterator[sqlite3.Connection]:
        connection = sqlite3.connect(self.db_path, timeout=30)
        connection.row_factory = sqlite3.Row
        connection.execute("PRAGMA foreign_keys = ON")
        connection.execute("PRAGMA journal_mode = WAL")
        try:
            yield connection
            connection.commit()
        except Exception:
            connection.rollback()
            raise
        finally:
            connection.close()

    def initialize(self) -> None:
        with self.connect() as connection:
            connection.executescript(
                """
                CREATE TABLE IF NOT EXISTS documents (
                    id INTEGER PRIMARY KEY,
                    path TEXT NOT NULL UNIQUE COLLATE NOCASE,
                    name TEXT NOT NULL,
                    suffix TEXT NOT NULL,
                    size INTEGER NOT NULL,
                    modified_at REAL NOT NULL,
                    indexed_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                    chunk_count INTEGER NOT NULL DEFAULT 0,
                    imported_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                    source_path TEXT NOT NULL DEFAULT '',
                    stored_path TEXT NOT NULL DEFAULT '',
                    sha256 TEXT NOT NULL DEFAULT '',
                    tags TEXT NOT NULL DEFAULT '',
                    category TEXT NOT NULL DEFAULT '未分类',
                    note TEXT NOT NULL DEFAULT '',
                    char_count INTEGER NOT NULL DEFAULT 0,
                    full_text TEXT NOT NULL DEFAULT ''
                );
                CREATE TABLE IF NOT EXISTS chunks (
                    id INTEGER PRIMARY KEY,
                    doc_id INTEGER NOT NULL REFERENCES documents(id) ON DELETE CASCADE,
                    location TEXT NOT NULL,
                    order_no INTEGER NOT NULL,
                    content TEXT NOT NULL
                );
                CREATE INDEX IF NOT EXISTS chunks_doc_id ON chunks(doc_id);
                CREATE TABLE IF NOT EXISTS bookmarks (
                    id INTEGER PRIMARY KEY,
                    chunk_id INTEGER NOT NULL UNIQUE REFERENCES chunks(id) ON DELETE CASCADE,
                    note TEXT NOT NULL DEFAULT '',
                    created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
                );
                CREATE TABLE IF NOT EXISTS pdf_annotations (
                    id INTEGER PRIMARY KEY,
                    doc_id INTEGER NOT NULL REFERENCES documents(id) ON DELETE CASCADE,
                    page_no INTEGER NOT NULL,
                    kind TEXT NOT NULL,
                    geometry TEXT NOT NULL,
                    color TEXT NOT NULL DEFAULT '#ffd54f',
                    text TEXT NOT NULL DEFAULT '',
                    created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                    updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
                );
                CREATE INDEX IF NOT EXISTS pdf_annotations_doc_page
                    ON pdf_annotations(doc_id, page_no);
                """
            )
            existing_columns = {
                str(row["name"]) for row in connection.execute("PRAGMA table_info(documents)").fetchall()
            }
            migrations = {
                "imported_at": "TEXT NOT NULL DEFAULT ''",
                "source_path": "TEXT NOT NULL DEFAULT ''",
                "stored_path": "TEXT NOT NULL DEFAULT ''",
                "sha256": "TEXT NOT NULL DEFAULT ''",
                "tags": "TEXT NOT NULL DEFAULT ''",
                "category": "TEXT NOT NULL DEFAULT '未分类'",
                "note": "TEXT NOT NULL DEFAULT ''",
                "char_count": "INTEGER NOT NULL DEFAULT 0",
                "full_text": "TEXT NOT NULL DEFAULT ''",
            }
            for column, definition in migrations.items():
                if column not in existing_columns:
                    connection.execute(f"ALTER TABLE documents ADD COLUMN {column} {definition}")
            connection.execute(
                """UPDATE documents SET
                   source_path=CASE WHEN source_path='' THEN path ELSE source_path END,
                   stored_path=CASE WHEN stored_path='' THEN path ELSE stored_path END,
                   category=CASE WHEN category='' THEN '未分类' ELSE category END,
                   imported_at=CASE WHEN imported_at='' THEN indexed_at ELSE imported_at END"""
            )
            try:
                connection.execute(
                    "CREATE VIRTUAL TABLE IF NOT EXISTS chunk_search USING fts5(content, tokenize='trigram')"
                )
                connection.executescript(
                    """
                    CREATE TRIGGER IF NOT EXISTS chunks_ai AFTER INSERT ON chunks BEGIN
                        INSERT INTO chunk_search(rowid, content) VALUES (new.id, new.content);
                    END;
                    CREATE TRIGGER IF NOT EXISTS chunks_ad AFTER DELETE ON chunks BEGIN
                        DELETE FROM chunk_search WHERE rowid = old.id;
                    END;
                    CREATE TRIGGER IF NOT EXISTS chunks_au AFTER UPDATE ON chunks BEGIN
                        DELETE FROM chunk_search WHERE rowid = old.id;
                        INSERT INTO chunk_search(rowid, content) VALUES (new.id, new.content);
                    END;
                    """
                )
                connection.execute(
                    """INSERT OR REPLACE INTO chunk_search(rowid, content)
                       SELECT id, content FROM chunks"""
                )
            except sqlite3.OperationalError:
                pass

    @staticmethod
    def file_hash(path: Path) -> str:
        digest = hashlib.sha256()
        with path.open("rb") as stream:
            for block in iter(lambda: stream.read(1024 * 1024), b""):
                digest.update(block)
        return digest.hexdigest()

    def duplicate_for(self, path: Path) -> sqlite3.Row | None:
        path = path.resolve()
        with self.connect() as connection:
            stored = connection.execute(
                "SELECT * FROM documents WHERE stored_path = ? COLLATE NOCASE LIMIT 1", (str(path),)
            ).fetchone()
            if stored:
                return stored
        digest = self.file_hash(path)
        with self.connect() as connection:
            return connection.execute(
                "SELECT * FROM documents WHERE name = ? COLLATE NOCASE AND sha256 = ? LIMIT 1",
                (path.name, digest),
            ).fetchone()

    def find_duplicates(self, paths: Iterable[str | Path]) -> list[Path]:
        duplicates: list[Path] = []
        for path in collect_supported_files(paths):
            if self.duplicate_for(path):
                duplicates.append(path)
        return duplicates

    @staticmethod
    def clean_category(category: str) -> str:
        cleaned = re.sub(r'[<>:"/\\|?*]', "_", category.strip()).strip(". ")
        return cleaned[:40] or "未分类"

    def _copy_to_storage(self, source: Path, digest: str, category: str) -> tuple[Path, bool]:
        safe_name = re.sub(r'[<>:"/\\|?*]', "_", source.name)
        category_dir = self.storage_dir / self.clean_category(category)
        category_dir.mkdir(parents=True, exist_ok=True)
        target = category_dir / f"{digest[:12]}_{safe_name}"
        existed = target.exists()
        if source.resolve() != target.resolve() and not existed:
            shutil.copy2(source, target)
        return target, not existed

    def import_file(
        self,
        path: Path,
        force: bool = False,
        duplicate_action: str = "skip",
        category: str = "未分类",
    ) -> tuple[str, int]:
        path = path.resolve()
        source_stat = path.stat()
        with self.connect() as connection:
            same_stored = connection.execute(
                "SELECT * FROM documents WHERE stored_path = ? COLLATE NOCASE LIMIT 1", (str(path),)
            ).fetchone()
            if same_stored and not force:
                stat_matches = (
                    int(same_stored["size"]) == source_stat.st_size
                    and abs(float(same_stored["modified_at"]) - source_stat.st_mtime) < 0.001
                )
                if stat_matches:
                    return "skipped", 0

        digest = self.file_hash(path)
        with self.connect() as connection:
            duplicate = connection.execute(
                "SELECT * FROM documents WHERE name = ? COLLATE NOCASE AND sha256 = ? LIMIT 1",
                (path.name, digest),
            ).fetchone()
            if duplicate and not force:
                # The content-addressed copy and index already exist. Re-reading
                # an identical open PDF/Word file only creates lock collisions.
                return "duplicate" if duplicate_action == "skip" else "skipped", 0

        category = self.clean_category(category)
        stored_path, created_copy = self._copy_to_storage(path, digest, category)
        stat = stored_path.stat()
        try:
            chunks = extract_file(stored_path)
        except Exception:
            if created_copy and stored_path.resolve() != path.resolve():
                stored_path.unlink(missing_ok=True)
            raise
        if not chunks:
            raise ValueError("文件中没有可建立索引的文字")
        full_text = "\n\n".join(chunk.content for chunk in chunks)
        char_count = len(re.sub(r"\s+", "", full_text))

        with self.connect() as connection:
            existing = None
            if duplicate and duplicate_action == "overwrite":
                existing = connection.execute("SELECT * FROM documents WHERE id = ?", (duplicate["id"],)).fetchone()
            if not existing:
                existing = connection.execute(
                    "SELECT * FROM documents WHERE stored_path = ? COLLATE NOCASE", (str(stored_path),)
                ).fetchone()
            if existing:
                doc_id = int(existing["id"])
                old_stored_path = Path(existing["stored_path"])
                connection.execute("DELETE FROM chunks WHERE doc_id = ?", (doc_id,))
                connection.execute(
                    """UPDATE documents SET path=?, name=?, suffix=?, size=?, modified_at=?,
                       indexed_at=CURRENT_TIMESTAMP, source_path=?, stored_path=?, sha256=?,
                       category=?, chunk_count=?, char_count=?, full_text=? WHERE id=?""",
                    (
                        str(stored_path), path.name, path.suffix.lower(), stat.st_size, stat.st_mtime,
                        str(path), str(stored_path), digest, category, len(chunks), char_count, full_text, doc_id,
                    ),
                )
                status = "updated"
            else:
                cursor = connection.execute(
                    """INSERT INTO documents(
                           path, name, suffix, size, modified_at, source_path, stored_path,
                           sha256, category, chunk_count, char_count, full_text, imported_at
                       ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                    (
                        str(stored_path), path.name, path.suffix.lower(), stat.st_size, stat.st_mtime,
                        str(path), str(stored_path), digest, category, len(chunks), char_count, full_text,
                        datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                    ),
                )
                doc_id = int(cursor.lastrowid)
                status = "added"
                old_stored_path = None
            connection.executemany(
                "INSERT INTO chunks(doc_id, location, order_no, content) VALUES (?, ?, ?, ?)",
                [(doc_id, chunk.location, chunk.order_no, chunk.content) for chunk in chunks],
            )
        if old_stored_path and old_stored_path != stored_path:
            try:
                if old_stored_path.resolve().is_relative_to(self.storage_dir.resolve()):
                    old_stored_path.unlink(missing_ok=True)
            except OSError:
                pass
        return status, len(chunks)

    def import_paths(
        self,
        paths: Iterable[str | Path],
        progress: Callable[[int, int, Path, str, str], None] | None = None,
        force: bool = False,
        duplicate_action: str = "skip",
        category: str = "未分类",
    ) -> dict[str, int | list[tuple[str, str]]]:
        files = collect_supported_files(paths)
        summary: dict[str, int | list[tuple[str, str]]] = {
            "total": len(files), "added": 0, "updated": 0, "skipped": 0,
            "duplicate": 0, "failed": []
        }
        for index, path in enumerate(files, 1):
            try:
                status, _count = self.import_file(
                    path, force=force, duplicate_action=duplicate_action, category=category
                )
                summary[status] = int(summary[status]) + 1
                message = status
            except PermissionError:
                failures = summary["failed"]
                assert isinstance(failures, list)
                failures.append((str(path), "文件正在被 Word、PDF 预览窗口或其他程序占用，请关闭后重试"))
                message = "failed"
            except Exception as exc:
                failures = summary["failed"]
                assert isinstance(failures, list)
                failures.append((str(path), str(exc)))
                message = "failed"
            if progress:
                progress(index, len(files), path, message, "")
        return summary

    def list_documents(self, category: str = "", sort_by: str = "name_asc") -> list[sqlite3.Row]:
        orders = {
            "name_asc": "name COLLATE NOCASE ASC",
            "name_desc": "name COLLATE NOCASE DESC",
            "date_desc": "imported_at DESC, name COLLATE NOCASE ASC",
            "date_asc": "imported_at ASC, name COLLATE NOCASE ASC",
        }
        order = orders.get(sort_by, orders["name_asc"])
        with self.connect() as connection:
            return connection.execute(
                f"SELECT * FROM documents WHERE (?='' OR category=?) ORDER BY {order}",
                (category, category),
            ).fetchall()

    def get_document(self, doc_id: int) -> sqlite3.Row | None:
        with self.connect() as connection:
            return connection.execute("SELECT * FROM documents WHERE id=?", (doc_id,)).fetchone()

    def get_document_by_path(self, path: str | Path) -> sqlite3.Row | None:
        resolved = str(Path(path).resolve())
        with self.connect() as connection:
            return connection.execute(
                """SELECT * FROM documents
                   WHERE path=? COLLATE NOCASE OR stored_path=? COLLATE NOCASE LIMIT 1""",
                (resolved, resolved),
            ).fetchone()

    def chunk_context(self, chunk_id: int, following: int = 2) -> list[sqlite3.Row]:
        """Return the matched chunk and nearby following chunks from the same document."""
        with self.connect() as connection:
            current = connection.execute(
                "SELECT doc_id, order_no FROM chunks WHERE id=?", (chunk_id,)
            ).fetchone()
            if not current:
                return []
            return connection.execute(
                """SELECT id, location, order_no, content FROM chunks
                   WHERE doc_id=? AND order_no>=? AND order_no<=?
                   ORDER BY order_no""",
                (current["doc_id"], current["order_no"], current["order_no"] + following),
            ).fetchall()

    def rename_document(self, doc_id: int, new_name: str) -> sqlite3.Row:
        clean_name = re.sub(r'[<>:"/\\|?*]', "_", Path(new_name.strip()).name).strip(". ")
        if not clean_name:
            raise ValueError("文件名不能为空")
        with self.connect() as connection:
            row = connection.execute("SELECT * FROM documents WHERE id=?", (doc_id,)).fetchone()
        if not row:
            raise ValueError("找不到该文档")
        old_path = Path(row["stored_path"])
        if not Path(clean_name).suffix:
            clean_name += old_path.suffix
        if Path(clean_name).suffix.lower() != old_path.suffix.lower():
            raise ValueError("重命名时不能改变文件类型")
        new_path = old_path.with_name(clean_name)
        if new_path.exists() and new_path.resolve() != old_path.resolve():
            raise FileExistsError("当前分类中已存在同名文件")
        old_path.rename(new_path)
        with self.connect() as connection:
            connection.execute(
                "UPDATE documents SET name=?, path=?, stored_path=? WHERE id=?",
                (clean_name, str(new_path), str(new_path), doc_id),
            )
        result = self.get_document(doc_id)
        assert result is not None
        return result

    def find_documents(self, keyword: str = "") -> list[sqlite3.Row]:
        with self.connect() as connection:
            return connection.execute(
                """SELECT * FROM documents
                   WHERE ?='' OR instr(lower(name), lower(?))>0
                      OR instr(lower(category), lower(?))>0
                      OR instr(lower(tags), lower(?))>0
                      OR instr(lower(note), lower(?))>0
                   ORDER BY name COLLATE NOCASE""",
                (keyword, keyword, keyword, keyword, keyword),
            ).fetchall()

    def counts(self) -> tuple[int, int]:
        with self.connect() as connection:
            docs = int(connection.execute("SELECT COUNT(*) FROM documents").fetchone()[0])
            chars = int(connection.execute("SELECT COALESCE(SUM(char_count), 0) FROM documents").fetchone()[0])
        return docs, chars

    def delete_documents(self, doc_ids: Iterable[int], delete_files: bool = True) -> int:
        ids = list(dict.fromkeys(int(value) for value in doc_ids))
        if not ids:
            return 0
        placeholders = ",".join("?" for _ in ids)
        with self.connect() as connection:
            paths = [Path(row[0]) for row in connection.execute(
                f"SELECT stored_path FROM documents WHERE id IN ({placeholders})", ids
            ).fetchall()]
            cursor = connection.execute(f"DELETE FROM documents WHERE id IN ({placeholders})", ids)
            deleted = cursor.rowcount
        if delete_files:
            storage_root = self.storage_dir.resolve()
            for path in paths:
                try:
                    resolved = path.resolve()
                    if resolved.is_relative_to(storage_root) and resolved.is_file():
                        resolved.unlink()
                except OSError:
                    pass
        return deleted

    def search(
        self, keyword: str, tag: str = "", category: str | Iterable[str] = "", doc_id: int = 0,
        limit: int = 5000,
    ) -> list[SearchHit]:
        if isinstance(category, str):
            categories = [category] if category else []
        else:
            categories = [str(item) for item in category if str(item).strip()]
        category_clause = ""
        if categories:
            category_clause = f" AND d.category IN ({','.join('?' for _ in categories)})"
        common_select = """SELECT c.id AS chunk_id, d.id AS doc_id, d.name, d.path,
                           c.location, c.content, d.modified_at"""
        if len(keyword) >= 3:
            sql = f"""
                {common_select}
                FROM chunk_search s
                JOIN chunks c ON c.id = s.rowid
                JOIN documents d ON d.id = c.doc_id
                WHERE s.content LIKE ?
                  AND (? = '' OR ',' || replace(d.tags, '，', ',') || ',' LIKE '%,' || ? || ',%')
                  {category_clause}
                  AND (? = 0 OR d.id = ?)
                LIMIT ?
            """
            params = (f"%{keyword}%", tag, tag, *categories, doc_id, doc_id, limit)
        else:
            sql = f"""
                {common_select}
                FROM chunks c JOIN documents d ON d.id = c.doc_id
                WHERE instr(lower(c.content), lower(?)) > 0
                  AND (? = '' OR ',' || replace(d.tags, '，', ',') || ',' LIKE '%,' || ? || ',%')
                  {category_clause}
                  AND (? = 0 OR d.id = ?)
                LIMIT ?
            """
            params = (keyword, tag, tag, *categories, doc_id, doc_id, limit)
        with self.connect() as connection:
            try:
                rows = connection.execute(sql, params).fetchall()
            except sqlite3.OperationalError:
                fallback = f"""
                    {common_select}
                    FROM chunks c JOIN documents d ON d.id = c.doc_id
                    WHERE instr(lower(c.content), lower(?)) > 0
                      AND (? = '' OR ',' || replace(d.tags, '，', ',') || ',' LIKE '%,' || ? || ',%')
                      {category_clause}
                      AND (? = 0 OR d.id = ?)
                    LIMIT ?
                """
                rows = connection.execute(
                    fallback, (keyword, tag, tag, *categories, doc_id, doc_id, limit)
                ).fetchall()
        hits = [
            SearchHit(
                chunk_id=int(row["chunk_id"]),
                doc_id=int(row["doc_id"]),
                name=str(row["name"]),
                path=str(row["path"]),
                location=str(row["location"]),
                content=str(row["content"]),
                snippet=make_snippet(str(row["content"]), keyword),
                modified_at=float(row["modified_at"]),
                score=str(row["content"]).casefold().count(keyword.casefold()),
            )
            for row in rows
        ]
        hits.sort(key=lambda hit: (-hit.score, hit.name.casefold(), hit.location))
        return hits

    def rebuild_all(self, progress=None) -> dict[str, int | list[tuple[str, str]]]:
        rows = self.list_documents()
        summary: dict[str, int | list[tuple[str, str]]] = {
            "total": len(rows), "added": 0, "updated": 0, "skipped": 0,
            "duplicate": 0, "failed": [],
        }
        for index, row in enumerate(rows, 1):
            path = Path(row["stored_path"])
            status = "updated"
            try:
                if not path.exists():
                    raise FileNotFoundError("资料库中的原始文件不存在")
                chunks = extract_file(path)
                full_text = "\n\n".join(chunk.content for chunk in chunks)
                char_count = len(re.sub(r"\s+", "", full_text))
                stat = path.stat()
                digest = self.file_hash(path)
                with self.connect() as connection:
                    connection.execute("DELETE FROM chunks WHERE doc_id=?", (row["id"],))
                    connection.executemany(
                        "INSERT INTO chunks(doc_id, location, order_no, content) VALUES (?, ?, ?, ?)",
                        [(row["id"], chunk.location, chunk.order_no, chunk.content) for chunk in chunks],
                    )
                    connection.execute(
                        """UPDATE documents SET size=?, modified_at=?, indexed_at=CURRENT_TIMESTAMP,
                           sha256=?, chunk_count=?, char_count=?, full_text=? WHERE id=?""",
                        (stat.st_size, stat.st_mtime, digest, len(chunks), char_count, full_text, row["id"]),
                    )
                summary["updated"] = int(summary["updated"]) + 1
            except Exception as exc:
                failures = summary["failed"]
                assert isinstance(failures, list)
                failures.append((str(path), str(exc)))
                status = "failed"
            if progress:
                progress(index, len(rows), path, status, "")
        return summary

    def all_tags(self) -> list[str]:
        tags: set[str] = set()
        for row in self.list_documents():
            tags.update(item.strip() for item in re.split(r"[,，]", str(row["tags"])) if item.strip())
        return sorted(tags, key=str.casefold)

    def list_categories(self) -> list[str]:
        categories = {"未分类"}
        categories.update(
            str(row["category"]) for row in self.list_documents() if str(row["category"]).strip()
        )
        categories.update(path.name for path in self.storage_dir.iterdir() if path.is_dir())
        return sorted(categories, key=lambda value: (value != "未分类", value.casefold()))

    def create_category(self, category: str) -> str:
        cleaned = self.clean_category(category)
        (self.storage_dir / cleaned).mkdir(parents=True, exist_ok=True)
        return cleaned

    def delete_category(self, category: str) -> int:
        """Remove a category without deleting documents; contents move to 未分类."""
        cleaned = self.clean_category(category)
        if cleaned == "未分类":
            raise ValueError("“未分类”是系统默认分类，不能删除")
        rows = self.list_documents(cleaned)
        for row in rows:
            self.move_document_category(int(row["id"]), "未分类")

        category_dir = (self.storage_dir / cleaned).resolve()
        storage_root = self.storage_dir.resolve()
        if category_dir.is_relative_to(storage_root) and category_dir.is_dir():
            fallback = self.storage_dir / "未分类"
            fallback.mkdir(parents=True, exist_ok=True)
            # Preserve any untracked files instead of silently deleting them.
            for item in list(category_dir.iterdir()):
                target = fallback / item.name
                if target.exists():
                    target = fallback / f"{cleaned}_{item.name}"
                counter = 2
                while target.exists():
                    target = fallback / f"{cleaned}_{counter}_{item.name}"
                    counter += 1
                shutil.move(str(item), str(target))
            category_dir.rmdir()
        return len(rows)

    def move_document_category(self, doc_id: int, category: str) -> None:
        category = self.create_category(category)
        with self.connect() as connection:
            row = connection.execute("SELECT * FROM documents WHERE id=?", (doc_id,)).fetchone()
        if not row:
            raise ValueError("找不到该文档")
        old_path = Path(row["stored_path"])
        if not old_path.exists():
            raise FileNotFoundError("学习库中的原始文档副本不存在")
        new_dir = self.storage_dir / category
        new_path = new_dir / old_path.name
        if old_path.resolve() != new_path.resolve():
            if new_path.exists():
                new_path = new_dir / f"{row['id']}_{old_path.name}"
            shutil.move(str(old_path), str(new_path))
        with self.connect() as connection:
            connection.execute(
                "UPDATE documents SET category=?, stored_path=?, path=? WHERE id=?",
                (category, str(new_path), str(new_path), doc_id),
            )

    def update_tags(self, doc_id: int, tags: str) -> None:
        cleaned = ",".join(dict.fromkeys(item.strip() for item in re.split(r"[,，]", tags) if item.strip()))
        with self.connect() as connection:
            connection.execute("UPDATE documents SET tags=? WHERE id=?", (cleaned, doc_id))

    def update_note(self, doc_id: int, note: str) -> None:
        with self.connect() as connection:
            connection.execute("UPDATE documents SET note=? WHERE id=?", (note.strip(), doc_id))

    def add_bookmark(self, chunk_id: int, note: str = "") -> None:
        with self.connect() as connection:
            connection.execute(
                """INSERT INTO bookmarks(chunk_id, note) VALUES (?, ?)
                   ON CONFLICT(chunk_id) DO UPDATE SET note=excluded.note""",
                (chunk_id, note.strip()),
            )

    def remove_bookmarks(self, bookmark_ids: Iterable[int]) -> int:
        ids = list(dict.fromkeys(int(value) for value in bookmark_ids))
        if not ids:
            return 0
        placeholders = ",".join("?" for _ in ids)
        with self.connect() as connection:
            cursor = connection.execute(f"DELETE FROM bookmarks WHERE id IN ({placeholders})", ids)
            return cursor.rowcount

    def list_bookmarks(self) -> list[sqlite3.Row]:
        with self.connect() as connection:
            return connection.execute(
                """SELECT b.id, b.note, b.created_at, c.id AS chunk_id, c.location,
                          c.content, d.id AS doc_id, d.name, d.path, d.category
                   FROM bookmarks b
                   JOIN chunks c ON c.id=b.chunk_id
                   JOIN documents d ON d.id=c.doc_id
                   ORDER BY b.created_at DESC"""
            ).fetchall()

    def checkpoint(self) -> None:
        with self.connect() as connection:
            connection.execute("PRAGMA wal_checkpoint(TRUNCATE)")

    def add_pdf_annotation(
        self, doc_id: int, page_no: int, kind: str, geometry: str,
        color: str = "#ffd54f", text: str = "",
    ) -> int:
        with self.connect() as connection:
            cursor = connection.execute(
                """INSERT INTO pdf_annotations(doc_id, page_no, kind, geometry, color, text)
                   VALUES (?, ?, ?, ?, ?, ?)""",
                (doc_id, page_no, kind, geometry, color, text.strip()),
            )
            return int(cursor.lastrowid)

    def list_pdf_annotations(self, doc_id: int, page_no: int | None = None) -> list[sqlite3.Row]:
        with self.connect() as connection:
            if page_no is None:
                return connection.execute(
                    "SELECT * FROM pdf_annotations WHERE doc_id=? ORDER BY page_no, id", (doc_id,)
                ).fetchall()
            return connection.execute(
                "SELECT * FROM pdf_annotations WHERE doc_id=? AND page_no=? ORDER BY id",
                (doc_id, page_no),
            ).fetchall()

    def update_pdf_annotation(self, annotation_id: int, color: str, text: str) -> None:
        with self.connect() as connection:
            connection.execute(
                """UPDATE pdf_annotations SET color=?, text=?, updated_at=CURRENT_TIMESTAMP
                   WHERE id=?""",
                (color, text.strip(), annotation_id),
            )

    def delete_pdf_annotation(self, annotation_id: int) -> None:
        with self.connect() as connection:
            connection.execute("DELETE FROM pdf_annotations WHERE id=?", (annotation_id,))

    def clear_pdf_annotations(self, doc_id: int) -> None:
        with self.connect() as connection:
            connection.execute("DELETE FROM pdf_annotations WHERE doc_id=?", (doc_id,))

    def export_backup(self, target: Path) -> None:
        documents = self.list_documents()
        manifest = {
            "version": 1,
            "created_at": datetime.now().isoformat(timespec="seconds"),
            "documents": [],
            "bookmarks": [],
            "pdf_annotations": [],
        }
        with zipfile.ZipFile(target, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for row in documents:
                stored = Path(row["stored_path"])
                archive_name = f"documents/{self.clean_category(row['category'])}/{row['id']}_{stored.name}"
                if stored.exists():
                    archive.write(stored, archive_name)
                text_name = f"texts/{row['id']}_{Path(row['name']).stem}.txt"
                archive.writestr(text_name, str(row["full_text"]).encode("utf-8"))
                manifest["documents"].append(
                    {
                        "name": row["name"], "archive_path": archive_name, "tags": row["tags"],
                        "category": row["category"], "note": row["note"], "sha256": row["sha256"],
                        "imported_at": row["imported_at"],
                    }
                )
            for bookmark in self.list_bookmarks():
                document = self.get_document(int(bookmark["doc_id"]))
                manifest["bookmarks"].append(
                    {
                        "document_sha256": document["sha256"] if document else "",
                        "document_name": bookmark["name"],
                        "location": bookmark["location"],
                        "note": bookmark["note"],
                    }
                )
            for document in documents:
                for annotation in self.list_pdf_annotations(int(document["id"])):
                    manifest["pdf_annotations"].append(
                        {
                            "document_sha256": document["sha256"],
                            "document_name": document["name"],
                            "page_no": annotation["page_no"],
                            "kind": annotation["kind"],
                            "geometry": annotation["geometry"],
                            "color": annotation["color"],
                            "text": annotation["text"],
                        }
                    )
            archive.writestr("manifest.json", json.dumps(manifest, ensure_ascii=False, indent=2).encode("utf-8"))

    def restore_backup(self, source: Path, progress=None) -> dict[str, int | list[tuple[str, str]]]:
        with zipfile.ZipFile(source, "r") as archive:
            try:
                manifest = json.loads(archive.read("manifest.json").decode("utf-8"))
            except Exception as exc:
                raise ValueError("备份包无效或 manifest.json 已损坏") from exc
            if manifest.get("version") != 1 or not isinstance(manifest.get("documents"), list):
                raise ValueError("备份包版本不受支持")
            with tempfile.TemporaryDirectory() as temp_dir:
                temp_root = Path(temp_dir)
                restored_items: list[tuple[Path, dict]] = []
                summary: dict[str, int | list[tuple[str, str]]] = {
                    "total": len(manifest["documents"]), "added": 0, "updated": 0,
                    "skipped": 0, "duplicate": 0, "failed": [],
                }
                for index, item in enumerate(manifest["documents"]):
                    archive_path = str(item.get("archive_path", ""))
                    if not archive_path.startswith("documents/") or ".." in Path(archive_path).parts:
                        raise ValueError("备份包包含不安全的文件路径")
                    filename = Path(str(item.get("name", ""))).name
                    if not filename:
                        continue
                    item_dir = temp_root / str(index)
                    item_dir.mkdir()
                    target = item_dir / filename
                    target.write_bytes(archive.read(archive_path))
                    restored_items.append((target, item))
                for index, (target, item) in enumerate(restored_items, 1):
                    status = "failed"
                    try:
                        status, _count = self.import_file(
                            target,
                            duplicate_action="overwrite",
                            category=str(item.get("category", "未分类")),
                        )
                        summary[status] = int(summary[status]) + 1
                        digest = self.file_hash(target)
                        with self.connect() as connection:
                            row = connection.execute(
                                "SELECT id FROM documents WHERE name=? COLLATE NOCASE AND sha256=?",
                                (target.name, digest),
                            ).fetchone()
                            if row:
                                connection.execute(
                                    "UPDATE documents SET tags=?, note=? WHERE id=?",
                                    (
                                        str(item.get("tags", "")), str(item.get("note", "")),
                                        row["id"],
                                    ),
                                )
                    except Exception as exc:
                        failures = summary["failed"]
                        assert isinstance(failures, list)
                        failures.append((str(target), str(exc)))
                    if progress:
                        progress(index, len(restored_items), target, status, "")
                with self.connect() as connection:
                    for bookmark in manifest.get("bookmarks", []):
                        row = connection.execute(
                            """SELECT c.id FROM chunks c JOIN documents d ON d.id=c.doc_id
                               WHERE d.name=? COLLATE NOCASE AND d.sha256=? AND c.location=? LIMIT 1""",
                            (
                                str(bookmark.get("document_name", "")),
                                str(bookmark.get("document_sha256", "")),
                                str(bookmark.get("location", "")),
                            ),
                        ).fetchone()
                        if row:
                            connection.execute(
                                """INSERT INTO bookmarks(chunk_id, note) VALUES (?, ?)
                                   ON CONFLICT(chunk_id) DO UPDATE SET note=excluded.note""",
                                (row["id"], str(bookmark.get("note", ""))),
                            )
                    for annotation in manifest.get("pdf_annotations", []):
                        document = connection.execute(
                            "SELECT id FROM documents WHERE name=? COLLATE NOCASE AND sha256=? LIMIT 1",
                            (
                                str(annotation.get("document_name", "")),
                                str(annotation.get("document_sha256", "")),
                            ),
                        ).fetchone()
                        if document:
                            connection.execute(
                                """INSERT INTO pdf_annotations(
                                       doc_id, page_no, kind, geometry, color, text
                                   ) VALUES (?, ?, ?, ?, ?, ?)""",
                                (
                                    document["id"], int(annotation.get("page_no", 0)),
                                    str(annotation.get("kind", "highlight")),
                                    str(annotation.get("geometry", "[]")),
                                    str(annotation.get("color", "#ffd54f")),
                                    str(annotation.get("text", "")),
                                ),
                            )
                return summary


def open_path(path: str) -> None:
    if os.name == "nt":
        os.startfile(path)
    else:
        subprocess.Popen(["xdg-open", path])


def reveal_path(path: str) -> None:
    if os.name == "nt":
        subprocess.Popen(["explorer", "/select,", str(Path(path).resolve())])
    else:
        open_path(str(Path(path).parent))


def open_at_location(path: str, location: str) -> None:
    source = Path(path).resolve()
    suffix = source.suffix.lower()
    number_match = re.search(r"(\d+)", location)
    number = int(number_match.group(1)) if number_match else 1

    if suffix == ".pdf":
        uri = source.as_uri() + f"#page={number}"
        edge_candidates = [
            shutil.which("msedge"),
            str(Path(os.getenv("PROGRAMFILES(X86)", "")) / "Microsoft/Edge/Application/msedge.exe"),
            str(Path(os.getenv("PROGRAMFILES", "")) / "Microsoft/Edge/Application/msedge.exe"),
        ]
        edge = next((item for item in edge_candidates if item and Path(item).exists()), None)
        if edge:
            subprocess.Popen([edge, uri])
        else:
            webbrowser.open(uri)
        return

    if suffix == ".docx":
        try:
            import pythoncom
            import win32com.client

            pythoncom.CoInitialize()
            word = win32com.client.Dispatch("Word.Application")
            word.Visible = True
            document = word.Documents.Open(str(source), ReadOnly=True, AddToRecentFiles=False)
            if location.startswith("表格") and document.Tables.Count >= number:
                document.Tables(number).Range.Select()
            elif document.Paragraphs.Count >= number:
                document.Paragraphs(number).Range.Select()
            word.Activate()
            pythoncom.CoUninitialize()
            return
        except Exception:
            pass

    if suffix in {".md", ".markdown"}:
        code = shutil.which("code")
        if code:
            subprocess.Popen([code, "-g", f"{source}:{number}"])
            return

    open_path(str(source))
